#%%
import pandas as pd
import numpy as np

import sklearn
from sklearn.decomposition import FactorAnalysis
from sklearn.decomposition import PCA

import matplotlib.pyplot as plt
import openpyxl as xl
import os
from sklearn import preprocessing

import statsmodels.api as sm

#%%
os.chdir("/Users/wonhyung64/data/redd/")
# os.chdir("E:\Data\\redd_plus\survey")
print("Current Working Directory :", os.getcwd())

#%%
data2 = pd.read_excel("redd_survey_case1_no_2.xlsx")
data3 = pd.read_excel("redd_survey_case1_no_3.xlsx")
#%%
# preprocessing

df_obsta = data2.iloc[:,:6]
df_value = data2.iloc[:,11:16]
df_knowl = data2.iloc[:,16:20]
df_recog = data3.iloc[:, 9:15]

df = pd.concat([df_obsta, df_value, df_knowl, df_recog],axis=1)
#%%
df = df.dropna(axis=0)
X1 = df.iloc[:,:6]
X2 = df.iloc[:,6:11]
X3 = df.iloc[:,11:15]
Y = df.iloc[:,15:]


scaler = preprocessing.StandardScaler().fit(Y)
Y_scaled = scaler.transform(Y)
Y_cols = ["Y1","Y2","Y3","Y4","Y5","Y6"]
Y_scaled = pd.DataFrame(Y_scaled, index=None, columns=Y_cols)
print(Y_scaled)
#%%

# factor analysis for all
X = pd.concat([X1,X2,X3],axis=1)
scaler = preprocessing.StandardScaler().fit(X)
X_scaled = scaler.transform(X)
#%%
feature_names = ['X1_1','X1_2','X1_3','X1_4','X1_5','X1_6','X2_1','X2_2','X2_3','X2_4','X2_5','X3_1','X3_2','X3_3','X3_4']
n_comps = 2
#%%
os.chdir("/Users/wonhyung64/data/redd")
# os.chdir("E:\Data\\redd_plus\\factor_analysis")
# %%

fa = FactorAnalysis(rotation = 'varimax')
fa.set_params(n_components = n_comps)
fa.fit(X_scaled)

loadings = fa.components_.T
col_tmp = ['loading1','loading2']
loading = pd.DataFrame(loadings, index=feature_names, columns = col_tmp)

variables = fa.fit_transform(X_scaled)
col_tmp = ['factor1','factor2']
factor = pd.DataFrame(variables, index=None, columns = col_tmp)

factor['intercept'] = 1

#%%
lm = sm.OLS(Y_scaled['Y1'], factor[['intercept','factor1','factor2']])
res = lm.fit()

summary = res.summary2()

loading.to_csv('y1_loading.csv', header = True, index=True)
factor.to_csv('y1_fator.csv', header = True, index=False)
summary.tables[0].to_csv('y1_summary1.csv', header = True, index=False)
summary.tables[1].to_csv('y1_summary2.csv', header = True, index=True)
summary.tables[2].to_csv('y1_summary3.csv', header = True, index=False)

# %%
lm = sm.OLS(Y_scaled['Y3'], factor[['intercept','factor1','factor2']])
res = lm.fit()

summary = res.summary2()

loading.to_csv('y3_loading.csv', header = True, index=True)
factor.to_csv('y3_fator.csv', header = True, index=False)
summary.tables[0].to_csv('y3_summary1.csv', header = True, index=False)
summary.tables[1].to_csv('y3_summary2.csv', header = True, index=True)
summary.tables[2].to_csv('y3_summary3.csv', header = True, index=False)

# %%
lm = sm.OLS(Y_scaled['Y5'], factor[['intercept','factor1','factor2']])
res = lm.fit()

summary = res.summary2()

loading.to_csv('y5_loading.csv', header = True, index=True)
factor.to_csv('y5_fator.csv', header = True, index=False)
summary.tables[0].to_csv('y5_summary1.csv', header = True, index=False)
summary.tables[1].to_csv('y5_summary2.csv', header = True, index=True)
summary.tables[2].to_csv('y5_summary3.csv', header = True, index=False)

# %%

lm = sm.OLS(Y_scaled['Y2'], factor[['intercept','factor1','factor2']])
res = lm.fit()

summary = res.summary2()

loading.to_csv('y2_loading.csv', header = True, index=True)
factor.to_csv('y2_fator.csv', header = True, index=False)
summary.tables[0].to_csv('y2_summary1.csv', header = True, index=False)
summary.tables[1].to_csv('y2_summary2.csv', header = True, index=True)
summary.tables[2].to_csv('y2_summary3.csv', header = True, index=False)
# %%
lm = sm.OLS(Y_scaled['Y4'], factor[['intercept','factor1','factor2']])
res = lm.fit()

summary = res.summary2()

loading.to_csv('y4_loading.csv', header = True, index=True)
factor.to_csv('y4_fator.csv', header = True, index=False)
summary.tables[0].to_csv('y4_summary1.csv', header = True, index=False)
summary.tables[1].to_csv('y4_summary2.csv', header = True, index=True)
summary.tables[2].to_csv('y4_summary3.csv', header = True, index=False)
# %%
lm = sm.OLS(Y_scaled['Y6'], factor[['intercept','factor1','factor2']])
res = lm.fit()

summary = res.summary2()

loading.to_csv('y6_loading.csv', header = True, index=True)
factor.to_csv('y6_fator.csv', header = True, index=False)
summary.tables[0].to_csv('y6_summary1.csv', header = True, index=False)
summary.tables[1].to_csv('y6_summary2.csv', header = True, index=True)
summary.tables[2].to_csv('y6_summary3.csv', header = True, index=False)

# %%

# %%

# %%

# %%

# %%

# %%

# %%
